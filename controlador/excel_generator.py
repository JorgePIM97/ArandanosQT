#controlador/excel_generator.py
import os
import pandas as pd
from PyQt6.QtWidgets import QMessageBox, QTableWidget
import win32api
import win32file
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

class ExcelGenerator:
    @staticmethod
    def tablewidget_to_dataframe(table: QTableWidget) -> pd.DataFrame:
        """Convierte QTableWidget a DataFrame con tipos de datos específicos"""
        try:
            # Obtener headers
            headers = []
            for col in range(table.columnCount()):
                header = table.horizontalHeaderItem(col)
                headers.append(header.text() if header else f"Columna {col+1}")
            
            # Obtener datos y convertir tipos
            data = []
            for row in range(table.rowCount()):
                row_data = []
                for col in range(table.columnCount()):
                    item = table.item(row, col)
                    text = item.text() if item else ""
                    
                    # Conversión de tipos según el nombre de la columna
                    header_text = headers[col].upper() if col < len(headers) else ""
                    
                    if "PESO" in header_text:
                        # Convertir a float, usar 0.0 si hay error
                        try:
                            row_data.append(float(text) if text else 0.0)
                        except:
                            row_data.append(0.0)
                    elif any(id_header in header_text for id_header in ["ID COSECHA", "ID RECOLECTOR", "ID ENTREGA"]):
                        # Convertir a entero, usar 0 si hay error
                        try:
                            row_data.append(int(text) if text else 0)
                        except:
                            row_data.append(0)
                    else:
                        # Mantener como texto
                        row_data.append(text)
                
                data.append(row_data)
            
            return pd.DataFrame(data, columns=headers)
        
        except Exception as e:
            raise Exception(f"Error al convertir tabla a DataFrame: {str(e)}")

    @staticmethod
    def save_to_excel(table: QTableWidget, parent_widget=None):
        """
        Guarda los datos de un QTableWidget en un archivo Excel en USB
        
        Args:
            table (QTableWidget): Tabla con los datos a exportar
            parent_widget (QWidget): Widget padre para mensajes
        """
        try:
            # Convertir tabla a DataFrame
            df = ExcelGenerator.tablewidget_to_dataframe(table)
            
            if df.empty:
                QMessageBox.information(parent_widget, "Información", "No hay datos para exportar.")
                return False

            # Detectar USB
            def get_usb_drives():
                drives = []
                for drive in win32api.GetLogicalDriveStrings().split('\x00')[:-1]:
                    if win32file.GetDriveType(drive) == 2:  # DRIVE_REMOVABLE (USB)
                        drives.append(drive)
                return drives

            usb_drives = get_usb_drives()
            
            if not usb_drives:
                QMessageBox.critical(parent_widget, "Error", "No se detectó ninguna memoria USB conectada.")
                return False

            # Generar nombre de archivo con fecha actual
            usb_path = usb_drives[0]
            fecha_actual = datetime.now().strftime('%Y%m%d_%H%M%S')
            file_name = f"CosechaArandano_{fecha_actual}.xlsx"
            full_path = os.path.join(usb_path, file_name)

            # Verificar si el archivo existe
            if os.path.exists(full_path):
                reply = QMessageBox.question(
                    parent_widget,
                    "Archivo existente",
                    f"El archivo {file_name} ya existe. ¿Sobrescribir?",
                    QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                )
                if reply == QMessageBox.StandardButton.No:
                    return False

            # Guardar Excel
            df.to_excel(full_path, index=False, engine='openpyxl')
            
            QMessageBox.information(
                parent_widget,
                "Éxito",
                f"Archivo guardado en:\n{full_path}\n\nRegistros exportados: {len(df)}"
            )
            return True

        except Exception as e:
            QMessageBox.critical(parent_widget, "Error", f"Error al exportar a Excel:\n{str(e)}")
            return False
        

class ExcelGeneratorSabana:

    def get_usb_drives(self):
        """Detecta unidades USB conectadas y retorna su ruta."""
        drives = []
        for drive in win32api.GetLogicalDriveStrings().split('\x00')[:-1]:
            # 2 = DRIVE_REMOVABLE
            if win32file.GetDriveType(drive) == 2:
                drives.append(drive)
        return drives

    def exportar_sabana_por_cuadrillero(self, df_sabana, fecha_inicio, fecha_fin, parent_widget=None):
        """
        Genera un archivo Excel por cada Cuadrillero en USB.
        """

        # Detectar USB
        usb_drives = self.get_usb_drives()

        if not usb_drives:
            QMessageBox.critical(parent_widget, "Error", "No se detectó ninguna memoria USB conectada.")
            return False

        # Usar la primera USB conectada
        usb_path = usb_drives[0]

        # Crear carpeta en la USB
        carpeta_salida = os.path.join(usb_path, "Reportes_Cosecha")
        if not os.path.exists(carpeta_salida):
            os.makedirs(carpeta_salida)

        # Obtener lista de cuadrilleros únicos
        cuadrilleros = df_sabana["Cuadrillero"].unique()

        for cuadrillero in cuadrilleros:

            df_filtrado = df_sabana[df_sabana["Cuadrillero"] == cuadrillero]

            # Obtener Clave
            clave = df_filtrado["Clave"].iloc[0]

            # Crear archivo Excel
            wb = Workbook()
            ws = wb.active
            ws.title = "Resumen"

            # -------------------------
            # ENCABEZADO
            # -------------------------
            ws["A1"] = f"Reporte de Cosecha del {fecha_inicio} al {fecha_fin}"
            ws["A1"].font = Font(size=14, bold=True)
            
            ws["A3"] = "Cuadrillero:"
            ws["B3"] = cuadrillero
            ws["A4"] = "Clave:"
            ws["B4"] = clave

            # -------------------------
            # TABLA DE DETALLES
            # -------------------------
            columnas = ["Nombre_Recolector", "Total_Vueltas", "Total_Peso", "Total_Buena", "Total_Regular", "Total_Mala"]

            start_row = 6
            for col_idx, col_name in enumerate(columnas, start=1):
                ws.cell(row=start_row, column=col_idx).value = col_name
                ws.cell(row=start_row, column=col_idx).font = Font(bold=True)

            # Insertar filas
            for i, row in df_filtrado.iterrows():
                for col_idx, col_name in enumerate(columnas, start=1):
                    ws.cell(row=start_row + 1 + i, column=col_idx).value = row[col_name]

            # -------------------------
            # TOTAL GENERAL
            # -------------------------
            total_vueltas = df_filtrado["Total_Vueltas"].sum()
            total_peso = df_filtrado["Total_Peso"].sum()

            total_row = start_row + len(df_filtrado) + 3

            ws[f"A{total_row}"] = "Total General:"
            ws[f"A{total_row}"].font = Font(bold=True)

            ws[f"B{total_row}"] = total_vueltas
            ws[f"B{total_row}"].font = Font(bold=True)

            ws[f"C{total_row}"] = f"{total_peso} Kg"
            ws[f"C{total_row}"].font = Font(bold=True)

            # Ajustar columnas
            for col in range(1, 10):
                ws.column_dimensions[chr(64 + col)].width = 20

            # -------------------------
            # GUARDAR ARCHIVO
            # -------------------------
            nombre_archivo = f"{cuadrillero.replace(' ', '_')}_{fecha_inicio}_{fecha_fin}.xlsx"
            ruta_final = os.path.join(carpeta_salida, nombre_archivo)
            wb.save(ruta_final)

        QMessageBox.information(
            parent_widget,
            "Éxito",
            f"Se generaron {len(cuadrilleros)} archivos en:\n{carpeta_salida}"
        )

        return True